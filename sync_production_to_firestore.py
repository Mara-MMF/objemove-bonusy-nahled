#!/usr/bin/env python3
"""Mirrors the production-data Google Sheet into Firestore, as raw grids, for
the --firebase-config public build to read at runtime instead of fetching the
sheet directly from a viewer's browser.

Run on a schedule by .github/workflows/sync-production.yml in the public repo
(https://github.com/Mara-MMF/objemove-bonusy-nahled) — not meant to be run by
hand except to test.

Env vars:
    FIREBASE_SERVICE_ACCOUNT_JSON  - the service-account key, as a JSON string
                                     (GitHub secret; written to this env var by
                                     the workflow, not a file on disk)
    SHEET_ID                       - defaults to the app's own SHEET_ID constant
    DIVISIONS_SHEET_ID             - workbook with the pojištění/investice/penze tabs
                                     (defaults to the app's DIVISIONS_SHEET_ID constant)

Requires: pip3 install requests openpyxl firebase-admin
"""
import io
import json
import os
import re
import sys

import requests
import openpyxl
import firebase_admin
from firebase_admin import credentials, firestore

SHEET_ID = os.environ.get("SHEET_ID", "1ZCFXSCFKr0zX0eMcifA5_f02J8AcaVIBl2KqqIS3gqo")
SHEET_TAB_NAME = "souhrn (objem)"
SHEET_TAB_NAME_KS = "souhrn (ks smluv)"
XLSX_EXPORT_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"

DIVISIONS_SHEET_ID = os.environ.get("DIVISIONS_SHEET_ID", "1V1HD9ykRJFjpfp8zWSZACC0LA1rrF3swpKW26Z6bPWE")
DIVISIONS_EXPORT_URL = f"https://docs.google.com/spreadsheets/d/{DIVISIONS_SHEET_ID}/export?format=xlsx"


def find_sheet_name(wb, wanted):
    if wanted in wb.sheetnames:
        return wanted
    wanted_norm = wanted.strip().lower()
    for name in wb.sheetnames:
        if name.strip().lower() == wanted_norm:
            return name
    return None


def sheet_to_grid(ws):
    """Same shape as the client's XLSX.utils.sheet_to_json(ws, {header:1, raw:true,
    defval:null}): a list of rows, each a list of cell values, None for empty cells.

    Trailing empty cells of every row are dropped (rows keep their index, even when empty).
    The client reads a missing cell exactly like an empty one, so nothing changes for it —
    but the souhrn (objem) grid shrinks from ~900 kB to ~120 kB, which matters because a
    Firestore document is capped at 1 MiB and that sheet gains columns every year."""
    grid = []
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        while cells and cells[-1] is None:
            cells.pop()
        grid.append(cells)
    return grid


def division_docs(xlsx_bytes):
    """Raw grids of every "Produkce <rok>" (Kč) and "Počet smluv <rok>" (ks) tab of the
    divisions workbook, keyed by the Firestore document id the app reads them from."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    docs = {}
    for name in wb.sheetnames:
        m = re.fullmatch(r"\s*Produkce\s+(\d{4})\s*", name)
        if m:
            docs[f"div_objem_{m.group(1)}"] = sheet_to_grid(wb[name])
            continue
        m = re.fullmatch(r"\s*Počet smluv\s+(\d{4})\s*", name)
        if m:
            docs[f"div_ks_{m.group(1)}"] = sheet_to_grid(wb[name])
    return docs


def main():
    key_json = os.environ.get("FIREBASE_SERVICE_ACCOUNT_JSON")
    if not key_json:
        raise SystemExit("Chybí env var FIREBASE_SERVICE_ACCOUNT_JSON.")

    resp = requests.get(XLSX_EXPORT_URL, timeout=60)
    resp.raise_for_status()

    wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)

    objem_name = find_sheet_name(wb, SHEET_TAB_NAME)
    if not objem_name:
        raise SystemExit(f'List "{SHEET_TAB_NAME}" nebyl v souboru nalezen.')
    objem_grid = sheet_to_grid(wb[objem_name])

    ks_name = find_sheet_name(wb, SHEET_TAB_NAME_KS)
    ks_grid = sheet_to_grid(wb[ks_name]) if ks_name else []

    cred = credentials.Certificate(json.loads(key_json))
    firebase_admin.initialize_app(cred)
    db = firestore.client()

    db.collection("production").document("objem").set({
        "json": json.dumps(objem_grid, ensure_ascii=False),
        "fetchedAt": firestore.SERVER_TIMESTAMP,
    })
    db.collection("production").document("ks").set({
        "json": json.dumps(ks_grid, ensure_ascii=False),
        "fetchedAt": firestore.SERVER_TIMESTAMP,
    })

    print(f"Nahráno do Firestore: objem {len(objem_grid)} řádků, ks {len(ks_grid)} řádků.")

    # The divisions workbook is a separate add-on: if it can't be read, the loans data
    # above is already synced — report the failure at the end so the run shows as failed.
    try:
        div_resp = requests.get(DIVISIONS_EXPORT_URL, timeout=60)
        div_resp.raise_for_status()
        docs = division_docs(div_resp.content)
        if not any(k.startswith("div_objem_") for k in docs):
            raise RuntimeError('V souboru divizí nebyl nalezen žádný list "Produkce <rok>".')
    except Exception as err:
        raise SystemExit(f"Divize (pojištění/investice/penze) se nepodařilo synchronizovat: {err}")
    for doc_id, grid in docs.items():
        db.collection("production").document(doc_id).set({
            "json": json.dumps(grid, ensure_ascii=False),
            "fetchedAt": firestore.SERVER_TIMESTAMP,
        })
    print("Nahráno do Firestore (divize): " + ", ".join(f"{k} ({len(v)} řádků)" for k, v in sorted(docs.items())))


if __name__ == "__main__":
    main()
